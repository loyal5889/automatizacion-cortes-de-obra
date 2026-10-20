"""Excel auditable con fórmulas visibles y caché calculada con Decimal."""
from datetime import datetime
from copy import copy
from decimal import Decimal
from pathlib import Path
import os
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.properties import CalcProperties

from .modelos import D, ErrorDatos

NAVY = "203C55"
BLUE = "1755A0"
GREEN = "23734B"
LIGHT = "EFF3F7"
YELLOW = "DCAD37"  # Amarillo predominante del logo de la empresa.
MONEY = '"$" #,##0.00;[Red]("$" #,##0.00);"$" 0.00'
QTY = '#,##0.00####;[Red](#,##0.00####);"–"'
RATE = '0.00%'
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def poner(ws, row, column, value):
    cell = ws.cell(row, column)
    cell.value = float(value) if isinstance(value, Decimal) else value
    # Los textos de fuentes no se ejecutan como fórmulas.
    if isinstance(value, str):
        cell.data_type = "s"
    cell.alignment = Alignment(vertical="center", wrap_text=True,
                               horizontal="right" if isinstance(value, (int, float, Decimal)) else "left")
    return cell


def formula(ws, address, expression, value, cache):
    ws[address] = expression
    cache.setdefault(ws.title, {})[address] = value
    ws[address].font = Font(name="Arial", size=10, color=GREEN if "!" in expression else "222222")
    ws[address].alignment = Alignment(vertical="center", wrap_text=True,
                                      horizontal="left" if isinstance(value, str) else "right")


def encabezado(ws, row, labels):
    for c, label in enumerate(labels, 1):
        cell = poner(ws, row, c, label)
        cell.font = Font(name="Arial", size=14, bold=True, color="222222")
        cell.fill = PatternFill("solid", fgColor=YELLOW)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(right=Side(style="thin", color="FFFFFF"))
    ws.row_dimensions[row].height = 70


def estilo(ws, widths, last_row):
    ws.sheet_view.showGridLines = False
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[col(i)].width = width * 1.4
    for row in ws.iter_rows(min_row=1, max_row=last_row, max_col=len(widths)):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="222222")
            cell.alignment = Alignment(vertical="center", wrap_text=True,
                                       horizontal="right" if isinstance(cell.value, (int, float)) else "left")
        ws.row_dimensions[row[0].row].height = 25
    # Imprimir al tamaño real: ajustar toda la tabla a una página reduce la letra.
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.scale = 100
    ws.page_setup.fitToWidth = None
    ws.page_setup.fitToHeight = None
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.oddFooter.center.text = "Página &P de &N"
    ws.print_area = f"A1:{col(len(widths))}{last_row}"


def titulo(ws, text):
    poner(ws, 2, 1, text)
    ws.merge_cells("A2:C2")
    ws["A2"].font = Font(name="Arial", size=14, bold=True, color="222222")
    for cell in ws[2][:3]:
        cell.fill = PatternFill("solid", fgColor=YELLOW)
    ws["A2"].alignment = Alignment(wrap_text=False)
    ws.row_dimensions[2].height = 29


def columnas_historicas(resultado):
    """Cantidades de cada corte: diferencia entre acumulados confirmados sucesivos."""
    anteriores = {a["item"]: D(0) for a in resultado.actividades}
    columnas = []
    for corte in resultado.cortes_anteriores:
        cantidades = {i: D(str(corte["cantidades_acumuladas"][str(i)])) for i in anteriores}
        delta = {i: cantidades[i] - anteriores[i] for i in anteriores}
        inicio = datetime.fromisoformat(corte["inicio"]).strftime("%d/%m/%Y")
        fin = datetime.fromisoformat(corte["fin"]).strftime("%d/%m/%Y")
        columnas.append((f"Cantidad corte {corte['numero']}\n{inicio} al {fin}", delta))
        anteriores = cantidades
    residual = {a["item"]: a["anterior"] - anteriores[a["item"]] for a in resultado.actividades}
    if any(residual.values()) or not columnas:
        etiqueta = "Cantidad anterior sin registro" if resultado.parametros.numero > 1 else "Cantidad anterior"
        columnas.append((etiqueta, residual))
    return columnas


def cachear(path, cache, sheet_names):
    """OOXML permite guardar fórmulas y resultados a la vez; openpyxl no expone cachés.

    Cada caché proviene del cálculo Decimal del mismo corte, no de resultados antiguos.
    Excel recalcula al abrir; esta caché permite leer con data_only sin tener Excel.
    """
    with ZipFile(path) as archive:
        files = {n: archive.read(n) for n in archive.namelist()}
    for index, name in enumerate(sheet_names, 1):
        key = f"xl/worksheets/sheet{index}.xml"
        root = ET.fromstring(files[key])
        for cell in root.iter(f"{{{NS}}}c"):
            address = cell.attrib.get("r")
            if address not in cache.get(name, {}):
                continue
            value = cache[name][address]
            tag = cell.find(f"{{{NS}}}v")
            if tag is None:
                tag = ET.SubElement(cell, f"{{{NS}}}v")
            if isinstance(value, str):
                cell.set("t", "str")
                tag.text = value
            else:
                cell.set("t", "n")
                tag.text = str(value)
        files[key] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)


def exportar(resultado, folder, definitivo=False):
    if definitivo and resultado.criticos:
        raise ErrorDatos("No se puede exportar como definitivo un corte con errores críticos.")
    folder.mkdir(parents=True, exist_ok=True)
    r, p, b = resultado, resultado.parametros, resultado.presupuesto
    generado = datetime.now().astimezone()
    status = "DEFINITIVO — verificar registro JSON del histórico" if definitivo else "BORRADOR — no registrado en el histórico"
    if r.criticos:
        status = "BORRADOR CON ERRORES — importes no aptos para cobro"
    wb = Workbook()
    wb.remove(wb.active)
    for name in ("Corte", "Consolidado reportes", "Presupuesto"):
        wb.create_sheet(name)
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    wb.properties.creator = "Cortes Darca II"
    wb.properties.title = f"Corte {p.numero:02d} — {r.config['proyecto']}"
    cache = {}

    ws = wb["Presupuesto"]
    estilo(ws, [8, 27, 80, 12, 18, 23, 24, 18], 30)
    titulo(ws, "Presupuesto contractual")
    poner(ws, 3, 1, f"Fuente: {b.archivo} / {r.config['hoja_presupuesto']}")
    ws["A3"].alignment = Alignment(wrap_text=False)
    encabezado(ws, 5, ["Ítem", "Concepto", "Descripción", "Unidad", "Cantidad contratada", "Precio unitario", "Valor contractual fuente", "Fila de origen"])
    budget_rows = {}
    for row, (i, item) in enumerate(sorted(b.items.items()), 6):
        budget_rows[i] = row
        for c, value in enumerate([i, item.concepto, item.descripcion, item.unidad, item.cantidad, item.precio, item.total, item.fila], 1):
            poner(ws, row, c, value)
        ws.row_dimensions[row].height = 58
        for c in (5, 6, 7):
            ws.cell(row, c).number_format = QTY if c == 5 else MONEY
        ws.cell(row, 6).number_format = MONEY
    rate_start = 6 + len(b.items) + 3
    budget_rates = {}
    for row, key in enumerate(("administracion", "imprevistos", "utilidad", "iva"), rate_start):
        budget_rates[key] = row
        poner(ws, row, 2, "IVA sobre utilidad" if key == "iva" else key.capitalize())
        poner(ws, row, 5, b.tasas[key]).number_format = RATE
    budget_total_row = rate_start + 6
    poner(ws, budget_total_row - 1, 2, "Costos directos contractuales")
    poner(ws, budget_total_row - 1, 7, b.directo).number_format = MONEY
    poner(ws, budget_total_row, 2, "Total contractual con AIU e IVA")
    poner(ws, budget_total_row, 7, b.total).number_format = MONEY
    ws.auto_filter.ref = f"A5:H{5 + len(b.items)}"
    ws.print_title_rows = "1:5"
    ws.print_area = f"A1:H{budget_total_row}"

    ws = wb["Consolidado reportes"]
    n = len(r.reportes)
    total_col, item_col, accumulated_col = 5, 7 + n, 6 + n
    last = 5 + len(r.consolidado)
    estilo(ws, [14, 57, 16, 12, 18] + [23] * n + [20, 15], last)
    titulo(ws, "Consolidado por WBS")
    poner(ws, 3, 1, "Fuente: DATOS. Las columnas semanales son cantidades ejecutadas; se incluyen semanas anteriores para calcular el acumulado.")
    ws["A3"].alignment = Alignment(wrap_text=False)
    encabezado(ws, 5, ["WBS", "Actividad", "Tramo", "Unidad", "Cantidad total del tramo"] +
                ["_\n".join(report.archivo.rsplit("_", 1)) for report in r.reportes] + ["Cantidad acumulada", "Ítem contractual"])
    ws.row_dimensions[5].height = 105
    for row, g in enumerate(r.consolidado, 6):
        for c, value in enumerate([g["wbs"], g["actividad"], g["tramo"], g["unidad"], g["total"]], 1):
            poner(ws, row, c, value)
        for c, report in enumerate(r.reportes, 6):
            # Ausencia no equivale a cero: queda vacía y se informa debajo de la tabla.
            poner(ws, row, c, g["semanas"].get(report.archivo))
        formula(ws, f"{col(accumulated_col)}{row}", f"=SUM(F{row}:{col(5+n)}{row})", g["acumulada"], cache)
        poner(ws, row, item_col, g["item"])
        for c in range(5, accumulated_col + 1):
            ws.cell(row, c).number_format = QTY
        ws.row_dimensions[row].height = 38
    ws.auto_filter.ref = f"A5:{col(item_col)}{last}"
    ws.print_title_rows = "1:5"

    ws = wb["Corte"]
    end = 8 + len(r.actividades)
    financial_start = end + 4
    historicas = columnas_historicas(r)
    actual_c = 6 + len(historicas)
    acum_c, precio_c, valor_c = actual_c + 1, actual_c + 2, actual_c + 3
    restante_c, saldo_c, porcentaje_c = actual_c + 4, actual_c + 5, actual_c + 6
    actual_l, acum_l, precio_l, valor_l = map(col, (actual_c, acum_c, precio_c, valor_c))
    restante_l, saldo_l, porcentaje_l = map(col, (restante_c, saldo_c, porcentaje_c))
    estilo(ws, [8, 35, 80, 11, 17] + [24] * len(historicas) + [17, 18, 22, 24, 17, 24, 16], financial_start + 20)
    # El resumen permanece en I aunque las columnas de la tabla se desplacen.
    ws.column_dimensions["I"].width = max(ws.column_dimensions["I"].width, 32 * 1.4)
    titulo(ws, f"{r.config['proyecto']} — Corte {p.numero:02d}")
    poner(ws, 3, 2, f"Periodo: {p.inicio:%d/%m/%Y} al {p.fin:%d/%m/%Y} · Valores en COP")
    ws["B3"].alignment = Alignment(wrap_text=False)
    poner(ws, 4, 2, status)
    ws["B4"].alignment = Alignment(wrap_text=False)
    ws["B4"].font = Font(name="Arial", size=10, bold=True, color="9C5700" if not definitivo else NAVY)
    encabezado(ws, 8, ["Ítem", "Actividad", "Descripción", "Unidad", "Cantidad contratada",  *[etiqueta for etiqueta, _ in historicas], f"Cantidad corte {p.numero} (actual)", "Cantidad ejecutada acumulada", "Precio unitario", "Valor corte actual", "Cantidad restante", "Valor restante", "Porcentaje ejecutado"])
    ws.row_dimensions[8].height = 86
    for blank in (5, 6, 7, end + 1, end + 2):
        ws.row_dimensions[blank].height = 12
    poner(ws, 5, 2, f"Generado: {generado:%d/%m/%Y %H:%M:%S} (UTC{generado:%z})").alignment = Alignment(wrap_text=False)
    ws.row_dimensions[5].height = 22
    # El logo original se escala proporcionalmente dentro de D1:L7.
    logo_path = Path(__file__).resolve().parents[1] / "assets" / "logo_empresa.png"
    logo = Image(str(logo_path))
    area_width = sum(int(ws.column_dimensions[col(c)].width * 7 + 5) for c in range(4, 13))
    area_height = sum(ws.row_dimensions[r].height * 96 / 72 for r in range(1, 8))
    scale = min((area_width - 16) / logo.width, (area_height - 16) / logo.height)
    width, height = logo.width * scale, logo.height * scale
    logo.anchor = OneCellAnchor(
        _from=AnchorMarker(col=3, row=0,
                           colOff=pixels_to_EMU((area_width - width) / 2),
                           rowOff=pixels_to_EMU((area_height - height) / 2)),
        ext=XDRPositiveSize2D(pixels_to_EMU(width), pixels_to_EMU(height)))
    ws.add_image(logo)
    for row, a in enumerate(r.actividades, 9):
        i = a["item"]
        item = b.items[i]
        br = budget_rows[i]
        poner(ws, row, 1, i)
        for target, source, value in (("B", "B", item.concepto), ("C", "C", item.descripcion),
                                      ("D", "D", item.unidad), ("E", "E", item.cantidad), (precio_l, "F", item.precio)):
            formula(ws, f"{target}{row}", f"='Presupuesto'!{source}{br}", value, cache)
        for cc, (_, cantidades) in enumerate(historicas, 6):
            poner(ws, row, cc, cantidades[i])
        sumexpr = f"SUMIFS('Consolidado reportes'!${col(accumulated_col)}$6:${col(accumulated_col)}${last},'Consolidado reportes'!${col(item_col)}$6:${col(item_col)}${last},A{row})"
        previous_expr = f"F{row}" if len(historicas) == 1 else f"SUM(F{row}:{col(actual_c-1)}{row})"
        formula(ws, f"{actual_l}{row}", f"={sumexpr}-{previous_expr}", a["actual"], cache)
        accumulated_expr = f"=F{row}+G{row}" if len(historicas) == 1 else f"=SUM(F{row}:{actual_l}{row})"
        formula(ws, f"{acum_l}{row}", accumulated_expr, a["acumulada"], cache)
        formula(ws, f"{valor_l}{row}", f"={actual_l}{row}*{precio_l}{row}", a["valor"], cache)
        formula(ws, f"{restante_l}{row}", f"=E{row}-{acum_l}{row}", a["restante"], cache)
        formula(ws, f"{saldo_l}{row}", f"={restante_l}{row}*{precio_l}{row}", a["restante"] * item.precio, cache)
        formula(ws, f"{porcentaje_l}{row}", f"={acum_l}{row}/E{row}", a["acumulada"] / item.cantidad, cache)
        for c in range(5, porcentaje_c):
            ws.cell(row, c).number_format = MONEY if c in (precio_c, valor_c, saldo_c) else QTY
        ws.cell(row, porcentaje_c).number_format = RATE
        ws.cell(row, 4).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 58
        if row % 2:
            for cell in ws[row][:porcentaje_c]:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
    f = r.finanzas
    names = ["directos", "administracion", "imprevistos", "utilidad", "iva", "bruto",
             "amortizacion", "retefuente", "garantia", "otros", "neto"]
    rows = {key: financial_start + idx for idx, key in enumerate(names)}
    labels = ["Costos directos ejecutados", "Administración", "Imprevistos", "Utilidad", "IVA sobre utilidad",
              "Valor bruto del corte", "Amortización del anticipo", "Retención en la fuente", "Retención de garantía", "Otros descuentos", "Valor neto a pagar"]
    poner(ws, financial_start - 1, 12, "Porcentaje").font = Font(name="Arial", size=10, bold=True, color=NAVY)
    poner(ws, financial_start - 1, 13, "Base").font = Font(name="Arial", size=10, bold=True, color=NAVY)
    exprs = {"directos": f"=SUM({valor_l}9:{valor_l}{end})", "administracion": f"=I{rows['directos']}*L{rows['administracion']}",
             "imprevistos": f"=I{rows['directos']}*L{rows['imprevistos']}", "utilidad": f"=I{rows['directos']}*L{rows['utilidad']}",
             "iva": f"=I{rows['utilidad']}*L{rows['iva']}", "bruto": f"=SUM(I{rows['directos']}:I{rows['iva']})",
             "neto": f"=I{rows['bruto']}-SUM(I{rows['amortizacion']}:I{rows['otros']})"}
    for key in ("amortizacion", "retefuente", "garantia"):
        rr = rows[key]
        base_expr = f'IF(M{rr}="bruto",I{rows["bruto"]},IF(M{rr}="directos",I{rows["directos"]},IF(M{rr}="utilidad",I{rows["utilidad"]},NA())))'
        if key == "retefuente":
            base_expr = f"(I{rows['bruto']}-I{rows['iva']})"
        exprs[key] = f"={base_expr}*L{rr}"
    for key, label in zip(names, labels):
        row = rows[key]
        poner(ws, row, 2, label)
        ws.cell(row, 2).alignment = Alignment(wrap_text=False)
        if key == "otros":
            poner(ws, row, 9, p.otros)
            ws.cell(row, 9).fill = PatternFill("solid", fgColor="FFF3CD")
        else:
            formula(ws, f"I{row}", exprs[key], f[key], cache)
        ws.cell(row, 9).number_format = MONEY
        if key in ("bruto", "neto"):
            for cell in ws[row][1:9]:
                cell.fill = PatternFill("solid", fgColor="DEE8EF")
                cell.font = Font(name="Arial", size=10, bold=True, color=NAVY)
        if key in ("administracion", "imprevistos", "utilidad", "iva", "amortizacion", "retefuente", "garantia"):
            rate = b.tasas[key] if key in b.tasas else getattr(p, key)
            if key in b.tasas:
                formula(ws, f"L{row}", f"='Presupuesto'!E{budget_rates[key]}", rate, cache)
            else:
                poner(ws, row, 12, rate).font = Font(name="Arial", size=10, color=BLUE)
                ws.cell(row, 12).fill = PatternFill("solid", fgColor="FFF3CD")
            ws.cell(row, 12).number_format = RATE
            if key in r.config["bases_deducciones"]:
                base_texto = "Bruto menos IVA" if key == "retefuente" else r.config["bases_deducciones"][key]
                poner(ws, row, 13, base_texto)
    row = financial_start + 12
    totals = [("Valor total contratado", f"='Presupuesto'!G{budget_total_row}", b.total),
              ("Valor ejecutado acumulado (AIU e IVA)", f"=SUMPRODUCT({acum_l}9:{acum_l}{end},{precio_l}9:{precio_l}{end})*(1+L{rows['administracion']}+L{rows['imprevistos']}+L{rows['utilidad']}*(1+L{rows['iva']}))", f["ejecutado_acumulado"]),
              ("Valor pendiente por ejecutar (AIU e IVA)", f"=I{row}-I{row+1}", f["pendiente"])]
    for rr, (label, expression, value) in enumerate(totals, row):
        poner(ws, rr, 2, label).alignment = Alignment(wrap_text=False)
        formula(ws, f"I{rr}", expression, value, cache)
        ws.cell(rr, 9).number_format = MONEY
    ws.conditional_formatting.add(f"{actual_l}9:{actual_l}{end}", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="FCE4D6")))
    ws.conditional_formatting.add(f"{porcentaje_l}9:{porcentaje_l}{end}", CellIsRule(operator="greaterThan", formula=["1"], fill=PatternFill("solid", fgColor="FCE4D6")))
    ws.auto_filter.ref = f"A8:{porcentaje_l}{end}"
    ws.print_title_rows = "1:8"
    ws.sheet_properties.tabColor = NAVY
    # La hoja principal muestra todos los resultados en un solo color sobrio.
    for row_cells in ws.iter_rows(min_row=9):
        for cell in row_cells:
            if cell.data_type == "f":
                cell.font = Font(name="Arial", size=10, color="222222", bold=cell.row in (rows["neto"], rows["bruto"]))

    metadata_start = financial_start + 17
    for rr, (label, value) in enumerate([
        ("Fecha inicial", p.inicio), ("Fecha final", p.fin),
        ("Número de corte", p.numero), ("Fecha y hora de generación", generado.replace(tzinfo=None, microsecond=0))
    ], metadata_start):
        poner(ws, rr, 2, label).alignment = Alignment(wrap_text=False)
        cell = poner(ws, rr, 9, value)
        if rr < metadata_start + 2:
            cell.number_format = "dd/mm/yyyy"
        elif rr == metadata_start + 3:
            cell.number_format = "dd/mm/yyyy hh:mm:ss"

    ws = wb["Consolidado reportes"]
    manifest_row = last + 4
    poner(ws, manifest_row - 1, 2, "Archivos procesados").font = Font(name="Arial", size=12, bold=True, color=NAVY)
    encabezado(ws, manifest_row, ["N.º", "Archivo", "Fecha Corte"])
    for rr, report in enumerate(r.reportes, manifest_row + 1):
        poner(ws, rr, 1, rr - manifest_row)
        poner(ws, rr, 2, report.archivo)
        poner(ws, rr, 3, report.fecha).number_format = "dd/mm/yyyy"
        ws.row_dimensions[rr].height = 30
    validation_row = manifest_row + len(r.reportes) + 4
    poner(ws, validation_row - 1, 2, "Validaciones y advertencias").font = Font(name="Arial", size=12, bold=True, color=NAVY)
    encabezado(ws, validation_row, ["Nivel", "Detalle", "Archivo", "", "", "", "", "", "", "Fila", "Ítem"])
    for rr, h in enumerate(r.hallazgos, validation_row + 1):
        poner(ws, rr, 1, h.nivel)
        poner(ws, rr, 2, f"{h.codigo}: {h.mensaje}")
        poner(ws, rr, 3, h.archivo).alignment = Alignment(vertical="center", wrap_text=False)
        poner(ws, rr, 10, h.fila)
        poner(ws, rr, 11, h.item)
        ws.row_dimensions[rr].height = max(44, 15 * ((len(h.mensaje) + len(h.codigo)) // 50 + 2))
    validation_end = validation_row + max(1, len(r.hallazgos))
    ws.conditional_formatting.add(f"A{validation_row+1}:A{validation_end}", CellIsRule(operator="equal", formula=['"ERROR"'], fill=PatternFill("solid", fgColor="FCE4D6")))
    ws.print_area = f"A1:{col(max(item_col, 11))}{validation_end}"
    for sheet in wb:
        sheet.freeze_panes = None
    # Títulos de secciones y del resumen: mismo amarillo y tipografía.
    for sheet, addresses in (
        (wb["Corte"], [f"L{financial_start - 1}", f"M{financial_start - 1}",
                       f"B{rows['bruto']}", f"B{rows['neto']}"]),
        (wb["Consolidado reportes"], [f"B{manifest_row - 1}", f"B{validation_row - 1}"]),
    ):
        for address in addresses:
            cell = sheet[address]
            cell.font = Font(name="Arial", size=14, bold=True, color="222222")
            cell.fill = PatternFill("solid", fgColor=YELLOW)
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            sheet.row_dimensions[cell.row].height = max(sheet.row_dimensions[cell.row].height or 25, 30)

    # Aplicar también a celdas creadas después del estilo inicial (avisos y metadatos).
    for sheet in wb:
        for cells in sheet:
            for cell in cells:
                font = copy(cell.font)
                font.sz = 20
                cell.font = font
            rr = cells[0].row
            if rr > 8 and any(cell.value is not None for cell in cells):
                sheet.row_dimensions[rr].height = max(35, (sheet.row_dimensions[rr].height or 25) * 1.4)
        sheet.oddFooter.center.size = 14

    base = f"CORTE_{p.numero:02d}_{r.config['codigo']}" + ("_DEFINITIVO" if definitivo else "")
    destination = folder / f"{base}_{generado:%Y%m%d_%H%M%S}.xlsx"
    if destination.exists():
        destination = folder / f"{base}_{generado:%Y%m%d_%H%M%S_%f}.xlsx"
    temp = None
    try:
        with tempfile.NamedTemporaryFile(dir=folder, suffix=".xlsx", delete=False) as handle:
            temp = Path(handle.name)
        wb.save(temp)
        cachear(temp, cache, wb.sheetnames)
        verificar_excel(temp, r)
        os.replace(temp, destination)
        return destination
    finally:
        wb.close()
        if temp is not None and temp.exists():
            temp.unlink()


def verificar_excel(path, resultado):
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    try:
        expected = ["Corte", "Consolidado reportes", "Presupuesto"]
        if values.sheetnames != expected:
            raise ErrorDatos("El Excel exportado no conserva sus hojas obligatorias.")
        if any(sheet.freeze_panes is not None for sheet in formulas):
            raise ErrorDatos("El Excel no debe contener filas ni columnas inmovilizadas.")
        ws = values["Corte"]
        historicas = columnas_historicas(resultado)
        actual_c = 6 + len(historicas)
        for row, a in enumerate(resultado.actividades, 9):
            total_anterior = sum((D(str(ws.cell(row, cc).value)) for cc in range(6, actual_c)), D(0))
            if abs(total_anterior - a["anterior"]) > D("0.00001"):
                raise ErrorDatos(f"El desglose de cortes anteriores no reconcilia en fila {row}.")
            for cc, (_, cantidades) in enumerate(historicas, 6):
                if abs(D(str(ws.cell(row, cc).value)) - cantidades[a["item"]]) > D("0.00001"):
                    raise ErrorDatos(f"Cantidad historica incorrecta en fila {row}.")
            for column, key in ((col(actual_c), "actual"), (col(actual_c+1), "acumulada"), (col(actual_c+3), "valor")):
                if abs(D(str(ws[f"{column}{row}"].value)) - a[key]) > D("0.00001"):
                    raise ErrorDatos(f"El Excel exportado no reconcilia en {column}{row}.")
        start = 12 + len(resultado.actividades)
        for offset, key in ((0, "directos"), (5, "bruto"), (10, "neto")):
            address = f"I{start + offset}"
            if abs(D(str(ws[address].value)) - resultado.finanzas[key]) > D("0.00001"):
                raise ErrorDatos(f"El Excel exportado no reconcilia {key}.")
        for sheet in formulas:
            for row in sheet:
                for cell in row:
                    cached = values[sheet.title][cell.coordinate]
                    if cached.data_type == "e" or (cell.data_type == "f" and cached.value is None):
                        raise ErrorDatos(f"Fórmula sin resultado o con error: {sheet.title}!{cell.coordinate}.")
        return True
    finally:
        formulas.close()
        values.close()
