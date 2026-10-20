"""Lectura exclusivamente de las hojas autorizadas; encabezados detectados por texto."""
import ast
from dataclasses import asdict
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from .modelos import (D, ErrorDatos, Hallazgo, Item, Presupuesto, Registro, Reporte,
                      cantidad, fecha, firma, hash_archivo, normalizar, numero, texto, unidad)


class Hoja:
    """Recalcula el subconjunto aritmético observado, sin ejecutar código de Excel."""

    def __init__(self, path, nombre):
        self.path = Path(path)
        try:
            self.wb = load_workbook(path, data_only=False)
        except Exception as exc:
            raise ErrorDatos(f"No se puede abrir {self.path.name}: {exc}") from exc
        if nombre not in self.wb.sheetnames:
            self.wb.close()
            raise ErrorDatos(f"{self.path.name}: falta la hoja obligatoria {nombre!r}.")
        self.ws = self.wb[nombre]
        self.cache = {}
        self.activos = set()

    def close(self):
        self.wb.close()

    def valor(self, coord):
        coord = coord.replace("$", "").upper()
        cell = self.ws[coord]
        if cell.data_type == "e":
            raise ErrorDatos(f"{self.path.name}!{coord}: error Excel {cell.value}.")
        if cell.data_type != "f":
            return cell.value
        if coord in self.cache:
            return self.cache[coord]
        if coord in self.activos:
            raise ErrorDatos(f"{self.path.name}!{coord}: fórmula circular.")
        self.activos.add(coord)
        try:
            expression = cell.value.lstrip("=").strip()

            def structured(match):
                table_name, header = match.groups()
                if table_name not in self.ws.tables:
                    raise ErrorDatos(f"Tabla desconocida {table_name} en {coord}.")
                table = self.ws.tables[table_name]
                left, top, right, bottom = range_boundaries(table.ref)
                if not top < cell.row <= bottom:
                    raise ErrorDatos(f"Referencia de tabla fuera de sus filas en {coord}.")
                cols = [c for c in range(left, right + 1)
                        if normalizar(self.ws.cell(top, c).value) == normalizar(header)]
                if len(cols) != 1:
                    raise ErrorDatos(f"Columna de tabla desconocida {header!r}.")
                return str(numero(self.valor(self.ws.cell(cell.row, cols[0]).coordinate)))

            expression = re.sub(r"([\w]+)\[\[#This Row\],\[([^\]]+)\]\]", structured, expression)

            def suma(match):
                left, top, right, bottom = range_boundaries(match.group(1).replace("$", ""))
                values = [self.valor(self.ws.cell(r, c).coordinate)
                          for r in range(top, bottom + 1) for c in range(left, right + 1)]
                return str(sum((numero(v) for v in values if v is not None and
                                not isinstance(v, str)), D(0)))

            expression = re.sub(r"SUM\(\s*(\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+)\s*\)", suma,
                                expression, flags=re.I)
            expression = re.sub(r"(?<![\w.])\$?[A-Z]{1,3}\$?\d+", lambda m: str(numero(self.valor(m.group(0)))),
                                expression, flags=re.I)
            tree = ast.parse(expression, mode="eval")

            def evaluar(node):
                if isinstance(node, ast.Expression):
                    return evaluar(node.body)
                if isinstance(node, ast.Constant) and type(node.value) in (int, float):
                    return numero(ast.get_source_segment(expression, node))
                if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
                    return evaluar(node.operand) * (-1 if isinstance(node.op, ast.USub) else 1)
                if isinstance(node, ast.BinOp):
                    a, b = evaluar(node.left), evaluar(node.right)
                    if isinstance(node.op, ast.Add): return a + b
                    if isinstance(node.op, ast.Sub): return a - b
                    if isinstance(node.op, ast.Mult): return a * b
                    if isinstance(node.op, ast.Div): return a / b
                raise ErrorDatos(f"Fórmula no admitida en {self.path.name}!{coord}: {cell.value}")

            value = evaluar(tree)
            if not value.is_finite():
                raise ErrorDatos(f"Fórmula no finita en {coord}.")
            self.cache[coord] = value
            return value
        except (SyntaxError, ArithmeticError, RecursionError) as exc:
            raise ErrorDatos(f"No se puede calcular {self.path.name}!{coord}: {cell.value}") from exc
        finally:
            self.activos.discard(coord)

    def encabezados(self, aliases):
        candidates = []
        for row in self.ws.iter_rows(min_row=1, max_row=self.ws.max_row):
            found = {}
            for cell in row:
                for key, names in aliases.items():
                    if normalizar(cell.value) in names:
                        if key in found:
                            raise ErrorDatos(f"{self.path.name}: encabezado duplicado {key} en fila {cell.row}.")
                        found[key] = cell.column
            if found:
                candidates.append((len(found), row[0].row, found))
        if not candidates:
            raise ErrorDatos(f"{self.path.name}: no se encontraron encabezados reconocibles.")
        _, row, found = max(candidates, key=lambda x: x[0])
        missing = set(aliases) - found.keys()
        if missing:
            raise ErrorDatos(f"{self.path.name}, fila {row}: columnas obligatorias faltantes: {', '.join(sorted(missing))}.")
        return row, found


def leer_config(root):
    try:
        config = json.loads((root / "config/proyecto.json").read_text(encoding="utf-8"))
        raw = json.loads((root / "config/mapeo_wbs.json").read_text(encoding="utf-8"))
        mapping = {}
        for key, values in raw.items():
            item = int(key)
            if item < 1 or not isinstance(values, list) or not values:
                raise ValueError("Cada ítem debe tener una lista no vacía de WBS")
            for wbs in values:
                wbs = texto(wbs)
                if not re.fullmatch(r"\d+(?:\.\d+)+", wbs) or wbs in mapping:
                    raise ValueError(f"WBS inválido o asignado dos veces: {wbs}")
                mapping[wbs] = item
        if set(config["bases_deducciones"]) != {"amortizacion", "retefuente", "garantia"}:
            raise ValueError("Faltan bases de deducciones")
        if any(v not in {"bruto", "directos", "utilidad"} for v in config["bases_deducciones"].values()):
            raise ValueError("Bases permitidas: bruto, directos, utilidad")
        if not re.fullmatch(r"[A-Za-z0-9_-]+", config["codigo"]):
            raise ValueError("Código de proyecto inválido")
        if type(config["items_esperados"]) is not int or config["items_esperados"] < 1:
            raise ValueError("Cantidad de ítems inválida")
        if not 0 <= config["decimales_cantidad"] <= 9:
            raise ValueError("Precisión de cantidades inválida")
        if numero(config["tolerancia_monetaria"]) < 0:
            raise ValueError("Tolerancia negativa")
        return config, mapping
    except (OSError, ValueError, KeyError, TypeError) as exc:
        raise ErrorDatos(f"Configuración inválida: {exc}") from exc


def leer_presupuesto(path, config, hallazgos):
    sheet = Hoja(path, config["hoja_presupuesto"])
    try:
        header, cols = sheet.encabezados({
            "item": {"item"}, "concepto": {"concepto"}, "descripcion": {"descripcion"},
            "unidad": {"unidad de medida", "unidad"}, "cantidad": {"cantidad"},
            "precio": {"valor unitario", "precio unitario"}, "total": {"valor total"}})
        # El archivo real rotula B4 como Item pero coloca los números en C.
        # Se verifica la columna numérica entre Item y Concepto, sin fijar C.
        candidates = []
        for col in range(cols["item"], cols["concepto"]):
            numbered = [(r, sheet.ws.cell(r, col).value) for r in range(header + 1, sheet.ws.max_row + 1)
                        if type(sheet.ws.cell(r, col).value) in (int, float)]
            if numbered:
                candidates.append((col, numbered))
        if len(candidates) != 1:
            raise ErrorDatos("No se identifica una única columna de números de ítem contractual.")
        items = {}
        for _, rows in candidates:
            for r, n in rows:
                if int(n) != n or n < 1 or int(n) in items:
                    raise ErrorDatos(f"Ítem contractual inválido o repetido en fila {r}: {n}.")
                n = int(n)
                vals = {key: sheet.valor(sheet.ws.cell(r, c).coordinate) for key, c in cols.items() if key != "item"}
                for key in ("concepto", "descripcion", "unidad"):
                    if not texto(vals[key]):
                        raise ErrorDatos(f"Presupuesto, fila {r}: falta {key}.")
                q, price, total = (numero(vals[k], f"Presupuesto fila {r}, {k}") for k in ("cantidad", "precio", "total"))
                if q <= 0 or price <= 0 or total <= 0:
                    raise ErrorDatos(f"Presupuesto fila {r}: cantidad, precio y total deben ser positivos.")
                if abs(q * price - total) > numero(config["tolerancia_monetaria"]):
                    hallazgos.append(Hallazgo("ERROR", "TOTAL_ITEM", f"Ítem {n}: cantidad × precio no coincide con valor total.", Path(path).name, r, n))
                items[n] = Item(n, texto(vals["concepto"]), texto(vals["descripcion"]), unidad(vals["unidad"]), q, price, total, r)
        if set(items) != set(range(1, config["items_esperados"] + 1)):
            raise ErrorDatos(f"Se esperaban ítems 1 a {config['items_esperados']}; encontrados {sorted(items)}.")
        labels = {"administracion": {"administracion"}, "imprevistos": {"imprevistos"},
                  "utilidad": {"utilidad"}, "iva": {"iva sobre utilidad"},
                  "total": {"valor total obra", "valort total obra", "valor total contractual"},
                  "directo": {"neto obra"}}
        values = {}
        for row in sheet.ws.iter_rows(min_row=header + 1):
            for cell in row:
                for key, names in labels.items():
                    if normalizar(cell.value) not in names:
                        continue
                    if key in values:
                        raise ErrorDatos(f"Etiqueta contractual duplicada: {key}.")
                    numeric = [sheet.valor(c.coordinate) for c in row[cell.column:] if c.value is not None]
                    numeric = [numero(v) for v in numeric if isinstance(v, (int, float, D))]
                    if key in {"total", "directo"}:
                        if len(numeric) != 1:
                            raise ErrorDatos(f"Valor contractual ambiguo o vacío para {key}.")
                        values[key] = numeric[0]
                    else:
                        rates = [v for v in numeric if 0 <= v <= 1]
                        if len(rates) != 1:
                            raise ErrorDatos(f"Porcentaje contractual ambiguo o vacío para {key}.")
                        values[key] = rates[0]
        if set(values) != set(labels):
            raise ErrorDatos(f"Faltan etiquetas del presupuesto: {sorted(set(labels) - set(values))}.")
        rates = {k: values[k] for k in ("administracion", "imprevistos", "utilidad", "iva")}
        direct = sum((i.cantidad * i.precio for i in items.values()), D(0))
        factor = 1 + rates["administracion"] + rates["imprevistos"] + rates["utilidad"] * (1 + rates["iva"])
        tol = numero(config["tolerancia_monetaria"])
        if abs(direct - values["directo"]) > tol or abs(direct * factor - values["total"]) > tol:
            hallazgos.append(Hallazgo("ERROR", "PRESUPUESTO_NO_RECONCILIA", "El total contractual no reconcilia con cantidades, precios, AIU e IVA.", Path(path).name))
        canonical = {"items": [asdict(i) for i in items.values()], "tasas": rates, "total": values["total"]}
        return Presupuesto(Path(path).name, items, rates, values["total"], direct, firma(canonical))
    finally:
        sheet.close()


def leer_reporte(path, config, hallazgos):
    sheet = Hoja(path, config["hoja_reportes"])
    try:
        header, cols = sheet.encabezados({name: {normalizar(name)} for name in
            ("Fecha Corte", "ID", "WBS", "Actividad", "Tramo", "Unidad", "Cantidad Total", "Cantidad Ejecutada Semana", "Observación")})
        records, seen = [], set()
        for r in range(header + 1, sheet.ws.max_row + 1):
            if all(sheet.ws.cell(r, c).value is None for c in cols.values()):
                continue
            vals = {key: sheet.valor(sheet.ws.cell(r, c).coordinate) for key, c in cols.items()}
            context = f"{Path(path).name}, fila {r}"
            for key in ("Fecha Corte", "ID", "WBS", "Tramo", "Unidad", "Cantidad Total", "Cantidad Ejecutada Semana"):
                if vals[key] is None or texto(vals[key]) == "":
                    raise ErrorDatos(f"{context}: falta {key}.")
            wbs = texto(vals["WBS"])
            if not re.fullmatch(r"\d+(?:\.\d+)+", wbs):
                raise ErrorDatos(f"{context}: WBS inválido {wbs!r}; se espera un código como 1.1.2.")
            if wbs in seen:
                raise ErrorDatos(f"{context}: WBS {wbs} repetido dentro del mismo reporte.")
            seen.add(wbs)
            total_raw = numero(vals["Cantidad Total"], context)
            qty_raw = numero(vals["Cantidad Ejecutada Semana"], context)
            if qty_raw < 0 or total_raw < 0:
                raise ErrorDatos(f"{context}: cantidades negativas.")
            total = cantidad(total_raw, config["decimales_cantidad"])
            qty = cantidad(qty_raw, config["decimales_cantidad"])
            if qty > total:
                hallazgos.append(Hallazgo("ERROR", "SEMANA_SUPERA_TOTAL", f"{wbs}: cantidad semanal superior al total del tramo.", Path(path).name, r))
            activity = texto(vals["Actividad"])
            if not activity:
                hallazgos.append(Hallazgo("ADVERTENCIA", "ACTIVIDAD_VACIA", f"WBS {wbs} sin nombre; se vincula por WBS.", Path(path).name, r))
            records.append(Registro(fecha(vals["Fecha Corte"]), texto(vals["ID"]), wbs, activity,
                                    texto(vals["Tramo"]), unidad(vals["Unidad"]), total, qty,
                                    texto(vals["Observación"]), Path(path).name, r))
        if not records:
            raise ErrorDatos(f"{Path(path).name}: no hay registros en DATOS.")
        if len({r.fecha for r in records}) != 1:
            raise ErrorDatos(f"{Path(path).name}: múltiples fechas de corte; se exige un reporte por semana.")
        canonical = sorted([(r.fecha.isoformat(), r.wbs, normalizar(r.tramo), r.unidad,
                            str(r.total), str(r.ejecutada)) for r in records])
        return Reporte(Path(path).name, records, firma(canonical), hash_archivo(Path(path)))
    finally:
        sheet.close()
