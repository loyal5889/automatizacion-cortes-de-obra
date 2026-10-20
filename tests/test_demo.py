import unittest, tempfile, shutil, sys
from pathlib import Path
from datetime import date

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from scripts.crear_demo import crear
from darca.calculo import preparar
from darca.modelos import Parametros,D
from darca.excel import exportar
from darca.historico import registrar
from openpyxl import load_workbook


class DemoTest(unittest.TestCase):
    def test_tres_cortes_y_retefuente(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);shutil.copytree(ROOT/'config',root/'config');crear(root)
            for n,start,end in [(1,5,10),(2,11,17),(3,18,24)]:
                r=preparar(root,Parametros(n,date(2026,1,start),date(2026,1,end),D('.1'),D('.02'),D('.1')))
                self.assertFalse(r.criticos)
                self.assertEqual(r.finanzas['retefuente'],(r.finanzas['bruto']-r.finanzas['iva'])*D('.02'))
                if n<3:registrar(root,r,exportar)
                else:
                    self.assertEqual(r.actividades[0]['actual'],D(30))
                    self.assertEqual(r.actividades[0]['acumulada'],D(90))
                    out=exportar(r,root/'output');w=load_workbook(out,data_only=True)
                    self.assertEqual([w['Corte'].cell(9,c).value for c in (6,7,8,9)],[20,40,30,90])
                    self.assertEqual(w.sheetnames,['Corte','Consolidado reportes','Presupuesto']);w.close()

    def test_demo_no_sobrescribe(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);shutil.copytree(ROOT/'config',root/'config');crear(root)
            with self.assertRaises(SystemExit):crear(root)
